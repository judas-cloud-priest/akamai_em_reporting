import os
import csv
import re
import time
import json
import logging
from datetime import datetime, timezone, timedelta

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from tqdm import tqdm
from tqdm.contrib.logging import logging_redirect_tqdm  # 💡 [추가] 로깅 충돌 방지 모듈

logger = logging.getLogger(__name__)

class PPTXOrchestrator:
    def __init__(self, template_path, output_path, api_client, global_args):
        self.template_path = template_path
        self.output_path = output_path
        self.api_client = api_client
        self.global_args = global_args
        
        try:
            self.prs = Presentation(self.template_path)
        except Exception as e:
            logger.error(f"Failed to load PPTX template: {e}")
            raise

    def _get_alt_text(self, shape):
        try:
            matches = shape.element.xpath('.//@descr')
            if matches:
                return matches[0]
        except Exception:
            pass
        return ""

    def _parse_slide_notes(self, slide, slide_idx):
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            text = slide.notes_slide.notes_text_frame.text
            if text:
                text = text.strip().replace('“', '"').replace('”', '"')
                
                # 💡 [핵심 추가] JSON 파싱 전 주석 제거 전처리 로직
                import re
                # 1. 멀티라인 주석 (/* ... */) 제거
                text = re.sub(r'/\*.*?\*/', '', text, flags=re.DOTALL)
                # 2. 한 줄 주석 (// ...) 제거 
                # (URL인 https:// 훼손 방지를 위해 줄 시작 부분에 있는 주석만 안전하게 제거)
                text = re.sub(r'^\s*//.*$', '', text, flags=re.MULTILINE)
                
                text = text.strip()
                
                if text.startswith('{'):
                    try:
                        return json.loads(text)
                    except json.JSONDecodeError as e:
                        logger.error(
                            f"\n❌ [Slide {slide_idx}] 슬라이드 노트(Global Params) JSON 구문 분석 실패!\n"
                            f"에러 원인: {e}\n"
                            f"문제가 발생한 JSON 원본:\n{text}\n"
                            f"{'-'*50}"
                        )
        return {}

    def _resolve_cpcode_list(self, cpcode_val):
        """'file:' 접두어가 있으면 CSV를 읽어오고, 아니면 원본 텍스트를 반환합니다."""
        if isinstance(cpcode_val, str) and cpcode_val.strip().startswith("file:"):
            # 'file:' 이후의 경로 추출 (앞뒤 공백 제거)
            file_path = cpcode_val.replace("file:", "", 1).strip()
            
            # 파이썬 스크립트 실행 위치(CWD)를 기준으로 절대 경로 생성
            abs_path = os.path.abspath(file_path)
            
            if not os.path.exists(abs_path):
                logger.error(f"❌ CPCode CSV 파일을 찾을 수 없습니다: {abs_path}")
                return ""
                
            cpcodes = []
            try:
                import csv
                # BOM(Byte Order Mark)이 포함된 UTF-8 엑셀 CSV도 안전하게 읽기 위해 utf-8-sig 사용
                with open(abs_path, 'r', encoding='utf-8-sig') as f:
                    reader = csv.reader(f)
                    for row in reader:
                        for cell in row:
                            cleaned = cell.strip()
                            # 순수 숫자(CP Code)만 추출하여 안전하게 배열에 담음
                            if cleaned.isdigit():
                                cpcodes.append(cleaned)
                
                logger.info(f"📁 CSV 파일에서 {len(cpcodes)}개의 CP Code를 성공적으로 로드했습니다. ({abs_path})")
            except Exception as e:
                logger.error(f"❌ CSV 읽기 실패 ({abs_path}): {e}")
                
            # 쉼표로 연결된 단일 문자열로 반환 (기존 하드코딩 방식과 완벽 호환)
            return ", ".join(cpcodes)
            
        return cpcode_val

    def _format_date_to_ymd(self, date_str):
        if not date_str: return ""
        try:
            dt_str = str(date_str).replace('Z', '+00:00')
            dt = datetime.fromisoformat(dt_str)
            return dt.strftime('%Y/%m/%d')
        except ValueError:
            return str(date_str)

    def _extract_month_abbr(self, date_str):
        if not date_str: return ""
        try:
            dt_str = str(date_str).replace('Z', '+00:00')
            return datetime.fromisoformat(dt_str).strftime('%b')
        except ValueError: return ""

    def _extract_year(self, date_str):
        if not date_str: return ""
        try:
            dt_str = str(date_str).replace('Z', '+00:00')
            return datetime.fromisoformat(dt_str).strftime('%Y')
        except ValueError: return ""

    def process_presentation(self):
        logger.info("Starting PPTX parsing and API orchestration...")
        
        self.current_start = self.global_args.start
        self.current_end = self.global_args.end
        self.current_cpcodes = getattr(self.global_args, 'cpcodes', None)
        self.current_service_name = getattr(self.global_args, 'service_name', None)
        
        # 💡 [핵심 추가] 에러가 발생한 슬라이드 번호를 추적하기 위한 집합(Set)
        self.failed_slides = set()
        
        target_slides = getattr(self.global_args, 'target_slides_list', [])
        if target_slides:
            logger.info(f"🎯 Targeted processing enabled for slides: {target_slides}")

        with logging_redirect_tqdm():
            for slide_idx, slide in tqdm(
                enumerate(self.prs.slides, start=1), 
                total=len(self.prs.slides), 
                desc="Processing Slides",
                bar_format="{l_bar}{bar}| {n_fmt}/{total_fmt} Slides [Elapsed: {elapsed}]"
            ):
                ignore_cache = self._should_ignore_cache_for_slide(slide_idx)
                
                # 1. 💡 상태(State) 동기화: 슬라이드 노트는 무조건 읽어서 전역 변수 업데이트 (전체 슬라이드 수행)
                slide_notes = self._parse_slide_notes(slide, slide_idx)
                if "global_params" in slide_notes:
                    params = slide_notes["global_params"]
                    self.current_start = params.get("START_DATE", self.current_start)
                    self.current_end = params.get("END_DATE", self.current_end)
                    
                    raw_cpcodes = params.get("CPCODE_LIST", self.current_cpcodes)
                    if hasattr(self, '_resolve_cpcode_list'):
                        self.current_cpcodes = self._resolve_cpcode_list(raw_cpcodes)
                    else:
                        self.current_cpcodes = raw_cpcodes
                        
                    self.current_service_name = params.get("SERVICE_NAME", self.current_service_name)
                
                # 2. 전역 텍스트 태그 설정
                global_tags = {}
                if self.current_service_name:
                    global_tags["{{SERVICE_NAME}}"] = str(self.current_service_name)
                if self.current_start:
                    global_tags["{{START_DATE}}"] = self._format_date_to_ymd(self.current_start)
                    global_tags["{{START_MONTH}}"] = self._extract_month_abbr(self.current_start)
                    global_tags["{{START_YEAR}}"] = self._extract_year(self.current_start)
                if self.current_end:
                    global_tags["{{END_DATE}}"] = self._format_date_to_ymd(self.current_end)
                    global_tags["{{END_MONTH}}"] = self._extract_month_abbr(self.current_end)
                    global_tags["{{END_YEAR}}"] = self._extract_year(self.current_end)
                
                # 3. 💡 타겟 슬라이드 여부 확인 (지정되지 않았으면 모두 True 처리)
                is_target_slide = (not target_slides) or (slide_idx in target_slides)

                # 💡 [핵심 추가] 타겟 슬라이드가 아니면 출력 PPTX 파일에서 '슬라이드 숨기기' 처리
                if not is_target_slide:
                    slide.element.set('show', '0')
                else:
                    # 만약 템플릿 원본에서 숨겨져 있던 슬라이드라면 타겟 지정 시 강제로 숨김 해제
                    if 'show' in slide.element.attrib:
                        del slide.element.attrib['show']

                for shape in slide.shapes:
                    # 텍스트 태그 치환은 모든 슬라이드에서 수행 (문서 일관성 유지)
                    if global_tags:
                        self._replace_tags_in_shape(shape, global_tags)
                    
                    # API 호출 및 도형 업데이트 로직은 타겟 슬라이드만 수행
                    if not is_target_slide:
                        continue
                    
                    alt_text = self._get_alt_text(shape)
                    if not alt_text:
                        continue
                        
                    # 💡 [핵심 추가] 도형이 표, 차트, 텍스트 박스 중 하나인지 판별
                    is_targetable_shape = getattr(shape, "has_table", False) or getattr(shape, "has_chart", False) or getattr(shape, "has_text_frame", False)
                    
                    if not alt_text.strip().startswith('{'):
                        # 데이터가 채워질 수 있는 도형인데 JSON이 아니라면 로그 출력 후 스킵
                        if is_targetable_shape:
                            preview = alt_text.strip()[:50].replace('\n', ' ')
                            logger.info(
                                f"⏭️ [Slide {slide_idx}] 표/차트/텍스트 박스의 대체 텍스트가 JSON 형식이 아니어서 일반 텍스트로 간주하고 스킵합니다. (preview: {preview}...)"
                            )
                        continue
                        
                    try:
                        alt_text = alt_text.replace('“', '"').replace('”', '"')
                        directive = json.loads(alt_text)
                    except json.JSONDecodeError as e:
                        self.failed_slides.add(slide_idx)
                        logger.error(
                            f"\n❌ [Slide {slide_idx}] 도형 대체 텍스트(Alt Text) JSON 구문 분석 실패!\n"
                            f"에러 원인: {e}\n"
                            f"문제가 발생한 JSON 원본:\n{alt_text}\n"
                            f"{'-'*50}"
                        )
                        continue
                        
                    self._process_shape_directive(slide_idx, shape, directive, ignore_cache)
                    time.sleep(0.15)
                
        self.prs.save(self.output_path)
        logger.info(f"Successfully processed and saved to {self.output_path}")
        
        # 💡 [핵심 추가] 스크립트 종료 시 실패한 슬라이드 목록 요약 및 가이드 출력
        if self.failed_slides:
            failed_str = ",".join(map(str, sorted(list(self.failed_slides))))
            logger.warning(
                f"\n{'='*70}\n"
                f"⚠️ [재실행 가이드] 일부 슬라이드의 처리가 완료되지 않았습니다 (API 실패/Timeout 등).\n"
                f"다음 파라미터를 복사하여 실패한 슬라이드만 빠르게 다시 실행해 보세요:\n\n"
                f"   --slides \"{failed_str}\"\n\n"
                f"{'='*70}"
            )
        else:
            logger.info("✅ 모든 슬라이드가 오류 없이 완벽하게 처리되었습니다!")

    def _align_date_for_interval(self, date_str, interval, is_end=False):
        if not date_str: return date_str
        try:
            # 파이썬 버전에 따른 타임존 파싱 버그를 원천 차단하기 위해 
            # 앞의 19자리(YYYY-MM-DDTHH:MM:SS)만 순수하게 잘라내서 계산합니다.
            clean_str = str(date_str)[:19] 
            dt = datetime.fromisoformat(clean_str)
            
            if interval.upper() == "DAY":
                if is_end:
                    if dt.hour != 0 or dt.minute != 0 or dt.second != 0:
                        dt = (dt + timedelta(days=1)).replace(hour=0, minute=0, second=0, microsecond=0)
                else:
                    dt = dt.replace(hour=0, minute=0, second=0, microsecond=0)
                    
            elif interval.upper() == "MONTH":
                # 다음 달 1일 00:00:00으로 완벽하게 스냅(Snap) 처리
                if is_end:
                    if dt.month == 12:
                        dt = dt.replace(year=dt.year+1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
                    else:
                        dt = dt.replace(month=dt.month+1, day=1, hour=0, minute=0, second=0, microsecond=0)
                else:
                    # 이번 달 1일 00:00:00으로 처리
                    dt = dt.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
                    
            # 계산 완료 후 Akamai가 요구하는 'Z'를 끝에 다시 붙여 반환
            return dt.strftime('%Y-%m-%dT%H:%M:%SZ')
        except Exception as e:
            logger.error(f"Date alignment failed for {date_str}: {e}")
            return str(date_str)

    def _process_shape_directive(self, slide_idx, shape, directive, ignore_cache):
        # 1. 템플릿에서 엔드포인트 가져오기
        endpoint_template = directive.get("endpoint", "")
        
        # 2. 전역 변수 설정
        api_start = self.current_start
        api_end = self.current_end
        active_service_name = self.current_service_name
        
        cpcode_list = []
        if self.current_cpcodes:
            cpcode_list = [c.strip() for c in str(self.current_cpcodes).split(',')]

        # 3. 날짜 정렬 적용
        if "interval=MONTH" in endpoint_template.upper():
            api_start = self._align_date_for_interval(api_start, "MONTH", is_end=False)
            api_end = self._align_date_for_interval(api_end, "MONTH", is_end=True)
        elif "interval=DAY" in endpoint_template.upper():
            api_start = self._align_date_for_interval(api_start, "DAY", is_end=False)
            api_end = self._align_date_for_interval(api_end, "DAY", is_end=True)

        # 4. 변수 치환
        endpoint = endpoint_template.replace("{START_DATE}", str(api_start)).replace("{END_DATE}", str(api_end))
        raw_payload = directive.get("payload", {})
        payload = self._replace_placeholders(raw_payload, cpcode_list, str(api_start), str(api_end), active_service_name)
        
        # 5. API 호출 및 렌더링
        try:
            response_data = self.api_client.request(
                method=directive.get("method", "POST"), 
                endpoint=endpoint, 
                payload=payload, 
                ignore_cache=ignore_cache
            )
            self._render_shape(slide_idx, shape, directive.get("action"), directive, response_data)
        except Exception as e:
            # 💡 [추가됨] API 에러 또는 차트 렌더링 실패 시 실패 슬라이드로 등록
            self.failed_slides.add(slide_idx)
            try:
                directive_dump = json.dumps(directive, indent=2, ensure_ascii=False)
            except Exception:
                directive_dump = str(directive)
                
            logger.error(
                f"\n❌ [Slide {slide_idx}] API 호출 또는 도형 업데이트 실패!\n"
                f"에러 원인: {e}\n"
                f"실행을 시도했던 JSON 지시어:\n{directive_dump}\n"
                f"{'-'*50}"
            )

    def _replace_placeholders(self, obj, cpcode_list, start_date, end_date, service_name=None):
        if isinstance(obj, dict):
            return {k: self._replace_placeholders(v, cpcode_list, start_date, end_date, service_name) for k, v in obj.items()}
        elif isinstance(obj, list):
            new_list = []
            for item in obj:
                if item == "{CPCODE_LIST}": new_list.extend(cpcode_list)
                else: new_list.append(self._replace_placeholders(item, cpcode_list, start_date, end_date, service_name))
            return new_list
        elif isinstance(obj, str):
            res = obj.replace("{START_DATE}", start_date).replace("{END_DATE}", end_date)
            if service_name: res = res.replace("{SERVICE_NAME}", str(service_name))
            return res
        return obj

    def _should_ignore_cache_for_slide(self, slide_index):
        if self.global_args.ignore_cache: return True
        if self.global_args.ignore_cache_slides:
            ignored_slides = [s.strip() for s in self.global_args.ignore_cache_slides.split(',')]
            if str(slide_index) in ignored_slides: return True
        return False

    def _render_shape(self, slide_idx, shape, action, directive, response_data):
        if action == "populate_table":
            self._handle_populate_table(slide_idx, shape, directive, response_data)
        elif action == "update_chart":
            self._handle_update_chart(slide_idx, shape, directive, response_data)
        elif action == "update_pie_chart":
            self._handle_update_pie_chart(slide_idx, shape, directive, response_data)
        elif action == "aggregate_response_classes":
            self._handle_aggregate_response_classes(slide_idx, shape, directive, response_data)
        elif action == "replace_text":
            self._handle_replace_text(slide_idx, shape, directive, response_data)
        else:
            logger.warning(f"[Slide {slide_idx}] Unknown action: {action}")

    def _format_x_axis_categories(self, categories):
        formatted = []
        last_date = None
        
        for cat in categories:
            if not cat:
                formatted.append("")
                continue
                
            dt, cat_str = None, str(cat).strip()
            
            try:
                if float(cat_str) > 1_000_000_000:
                    dt = datetime.fromtimestamp(float(cat_str), tz=timezone.utc)
            except ValueError:
                pass
                
            if dt is None:
                try:
                    dt = datetime.fromisoformat(cat_str.replace('Z', '+00:00'))
                except ValueError:
                    pass
            
            if dt is not None:
                current_date = dt.strftime('%m.%d')
                
                # 💡 [핵심 로직] 날짜가 바뀌었을 때(하루의 첫 데이터)만 라벨을 출력합니다.
                if last_date != current_date:
                    formatted.append(current_date)
                    last_date = current_date
                else:
                    # 같은 날짜의 나머지 23시간은 빈칸("")으로 처리하여 X축 중복 출력을 방지합니다.
                    formatted.append("")
            else:
                formatted.append(cat_str)
                
        return formatted

    def _apply_unit_formatting(self, value, target_tag, metric=""):
        target_upper = target_tag.upper()
        metric_upper = metric.upper()
        
        # 1. CP Code 처리 (앞자리 '0' 보존)
        if "CPCODE" in target_upper or "CPCODE" in metric_upper:
            val_str = str(value)
            if val_str.endswith('.0'):
                return val_str[:-2]
            return val_str

        # 2. 순수 텍스트(문자열) 감지 및 원본 반환
        try:
            float_val = float(value)
        except (ValueError, TypeError):
            return str(value)

        # 3. 퍼센트(%) 처리
        if "PERCENT" in target_upper or "PCT" in target_upper or "ACC_VOL" in target_upper or "PERCENT" in metric_upper or "OFFLOAD" in target_upper:
            return f"{float_val:,.2f}%"
            
        # 4. 강제 카운트(건수) 처리
        if "CNT" in target_upper or "COUNT" in target_upper or "ERR" in target_upper:
            return f"{float_val:,.0f}"
            
        # 5. 단위 식별 키워드 설정 (BYTE 포함 추가)
        is_hits = "HITS" in target_upper or "HPS" in target_upper or "HITS" in metric_upper
        is_bytes = "BYTE" in target_upper or "BYTE" in metric_upper
        is_bps = "BPS" in target_upper or "BITSPERSECOND" in metric_upper
        
        # 6. 각 단위별 접미사(Suffix) 정의 - 💡 [핵심 수정] BYTES 조건을 BPS보다 먼저 확인합니다!
        if is_hits:
            if "PEAK" in target_upper or "HPS" in target_upper:
                u_p, u_t, u_g, u_m, u_k, u_base = "P Hits/s", "T Hits/s", "G Hits/s", "M Hits/s", "K Hits/s", "Hits/s"
            else:
                u_p, u_t, u_g, u_m, u_k, u_base = "P Hits", "T Hits", "G Hits", "M Hits", "K Hits", "Hits"
        elif is_bytes:  # <--- BYTES가 우선합니다! (EDGE_BYTES 등은 무조건 이쪽으로 빠짐)
            u_p, u_t, u_g, u_m, u_k, u_base = "PB", "TB", "GB", "MB", "KB", "B"
        elif is_bps:
            u_p, u_t, u_g, u_m, u_k, u_base = "Pbps", "Tbps", "Gbps", "Mbps", "Kbps", "bps"
        else:
            # Bytes, Hits, bps 어느 것에도 해당하지 않는 '일반 숫자' 처리
            if float_val.is_integer():
                return f"{float_val:,.0f}"
            else:
                return f"{float_val:,.2f}"
            
        # 7. 지정된 단위가 있는 경우에만 스케일링 적용 (K, M, G, T, P)
        if float_val >= 1_000_000_000_000_000: return f"{float_val / 1_000_000_000_000_000:,.2f} {u_p}"
        elif float_val >= 1_000_000_000_000: return f"{float_val / 1_000_000_000_000:,.2f} {u_t}"
        elif float_val >= 1_000_000_000: return f"{float_val / 1_000_000_000:,.2f} {u_g}"
        elif float_val >= 1_000_000: return f"{float_val / 1_000_000:,.2f} {u_m}"
        elif float_val >= 1_000: return f"{float_val / 1_000:,.2f} {u_k}"
        else: return f"{float_val:,.0f} {u_base}"

    def _apply_post_processing(self, data_rows, rules):
        if not rules: return data_rows
        
        # 1. Row-level calculate
        for calc in rules.get("calculate", []):
            if calc.get("operation") == "sum" and calc.get("new_field"):
                for row in data_rows:
                    row[calc["new_field"]] = sum(float(row.get(f, 0.0)) for f in calc.get("fields", []) if row.get(f) is not None)

        # 2. Filters
        filters = rules.get("filters")
        if filters:
            logic, conditions = filters.get("logic", "AND").upper(), filters.get("conditions", [])
            filtered_rows = []
            for row in data_rows:
                match_results = []
                for cond in conditions:
                    row_val = row.get(cond.get("field"))
                    if row_val is None:
                        match_results.append(False); continue
                    try:
                        op, val, pat = cond.get("operator"), cond.get("value"), cond.get("pattern", "")
                        if op == "<=": res = float(row_val) <= float(val)
                        elif op == ">=": res = float(row_val) >= float(val)
                        elif op == "<": res = float(row_val) < float(val)
                        elif op == ">": res = float(row_val) > float(val)
                        elif op == "==": res = str(row_val) == str(val)
                        elif op == "regex_match": res = bool(re.search(pat, str(row_val), re.IGNORECASE))
                        elif op == "regex_not_match": res = not bool(re.search(pat, str(row_val), re.IGNORECASE))
                        match_results.append(res)
                    except ValueError: match_results.append(False)
                if not conditions: filtered_rows.append(row)
                elif logic == "AND" and all(match_results): filtered_rows.append(row)
                elif logic == "OR" and any(match_results): filtered_rows.append(row)
            data_rows = filtered_rows

        # 3. Group By
        group_by = rules.get("group_by")
        if group_by:
            key, metrics = group_by.get("key"), group_by.get("metrics", [])
            grouped = {}
            for row in data_rows:
                grp_val = str(row.get(key, ""))
                if grp_val not in grouped:
                    grouped[grp_val] = {key: grp_val}
                    for m in metrics: grouped[grp_val][m] = 0.0
                for m in metrics: grouped[grp_val][m] += float(row.get(m, 0.0))
            data_rows = list(grouped.values())

        # 4. Percentage 계산
        calc_metrics = rules.get("calculate_metrics", [])
        for calc in calc_metrics:
            if calc.get("type") == "percentage":
                src, new_f = calc.get("source_field"), calc.get("new_field")
                total_sum = sum(float(r.get(src, 0.0)) for r in data_rows)
                for r in data_rows: r[new_f] = (float(r.get(src, 0.0)) / total_sum * 100) if total_sum > 0 else 0.0

        # 5. Sort
        sort_rule = rules.get("sort")
        if sort_rule:
            field, order = sort_rule.get("field"), sort_rule.get("order", "DESC").upper()
            def sort_key(x):
                try: return float(x.get(field, 0))
                except (ValueError, TypeError): return str(x.get(field, 0)).lower()
            data_rows = sorted(data_rows, key=sort_key, reverse=(order == "DESC"))

        # 6. Cumulative 계산
        for calc in calc_metrics:
            if calc.get("type") == "cumulative":
                src, new_f, acc = calc.get("source_field"), calc.get("new_field"), 0.0
                for r in data_rows:
                    acc += float(r.get(src, 0.0))
                    r[new_f] = acc

        # 7. Limit
        if rules.get("limit"): data_rows = data_rows[:int(rules.get("limit"))]
        return data_rows

    def _replace_tags_in_shape(self, shape, tag_values):
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            for tag, val in tag_values.items():
                                if tag in run.text: run.text = run.text.replace(tag, str(val))
        elif shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    for tag, val in tag_values.items():
                        if tag in run.text: run.text = run.text.replace(tag, str(val))

    def _handle_populate_table(self, slide_idx, shape, directive, response_data):
        extract_array = directive.get("extract_array", [])
        if not extract_array: return
        
        processed_rows = self._apply_post_processing(response_data.get("data", []), directive.get("post_processing", {}))
        tag_values = {}
        for idx, row in enumerate(processed_rows, start=1):
            for rule in extract_array:
                metric = rule.get("metric", "")
                raw_val = row.get(metric, 0.0)
                try: calc_val = float(raw_val) * rule.get("multiplier", 1.0)
                except: calc_val = raw_val
                target_tag = f"{{{{{rule.get('tag_prefix')}{idx}}}}}"
                tag_values[target_tag] = calc_val if isinstance(calc_val, str) else self._apply_unit_formatting(calc_val, target_tag, metric)

        limit = directive.get("post_processing", {}).get("limit", len(processed_rows))
        for missing_idx in range(len(processed_rows) + 1, limit + 1):
            for rule in extract_array:
                tag_values[f"{{{{{rule.get('tag_prefix')}{missing_idx}}}}}"] = "-" 
        self._replace_tags_in_shape(shape, tag_values)

    def _handle_update_chart(self, slide_idx, shape, directive, response_data):
        if not shape.has_chart: return
        mapping = directive.get("mapping", {})
        x_axis_key, series_mappings = mapping.get("x_axis"), mapping.get("series", [])
        data_rows = response_data.get("data", [])
        
        raw_categories = []
        for row in data_rows:
            cat = row.get(x_axis_key, "")
            if cat not in raw_categories: raw_categories.append(cat)
                
        categories = self._format_x_axis_categories(raw_categories)
        raw_series_data, global_max = {}, 0.0
        
        for s_map in series_mappings:
            s_name, s_metric, s_cond = s_map.get("name"), s_map.get("metric"), s_map.get("condition")
            s_values = []
            for x_val in raw_categories:
                val_sum = 0.0
                for row in data_rows:
                    if row.get(x_axis_key, "") == x_val:
                        if s_cond:
                            if str(row.get(s_cond.get("field"), "")) == str(s_cond.get("value")): val_sum += float(row.get(s_metric, 0))
                        else: val_sum += float(row.get(s_metric, 0))
                s_values.append(val_sum)
            raw_series_data[s_name] = s_values
            if s_values: global_max = max(global_max, max(s_values))

        scale_factor, unit_suffix = 1, ""
        if global_max >= 1_000_000_000: scale_factor, unit_suffix = 1_000_000_000, " (G)"
        elif global_max >= 1_000_000: scale_factor, unit_suffix = 1_000_000, " (M)"
        elif global_max >= 1_000: scale_factor, unit_suffix = 1_000, " (K)"

        chart_data = CategoryChartData()
        chart_data.categories = categories
        for s_name, s_values in raw_series_data.items():
            chart_data.add_series(f"{s_name}{unit_suffix}", tuple([v / scale_factor for v in s_values]))
        shape.chart.replace_data(chart_data)

        # 👇 [여기에 추가] 범례(Legend) 강제 새로고침 트릭
        if shape.chart.has_legend:
            shape.chart.has_legend = False
            shape.chart.has_legend = True
            
        logger.info(f"[Slide {slide_idx}] Chart dynamically updated with {len(categories)} categories and {len(series_mappings)} series.")

    def _handle_update_pie_chart(self, slide_idx, shape, directive, response_data):
        if not shape.has_chart: return
        mapping = directive.get("mapping", {})
        category_key = mapping.get("category", "country")
        value_key = mapping.get("value", "edgeBytes")
        
        aggregated_data = {}
        for row in response_data.get("data", []):
            cat = str(row.get(category_key, "")).strip() or "Unknown"
            aggregated_data[cat] = aggregated_data.get(cat, 0.0) + float(row.get(value_key, 0.0))

        sorted_items = sorted(aggregated_data.items(), key=lambda x: x[1], reverse=True)
        limit = directive.get("post_processing", {}).get("limit")
        if limit and len(sorted_items) > int(limit):
            top_items = sorted_items[:int(limit)]
            others_sum = sum(item[1] for item in sorted_items[int(limit):])
            if others_sum > 0: top_items.append(("Others", others_sum))
            sorted_items = top_items

        if not sorted_items: return
        chart_data = CategoryChartData()
        chart_data.categories = [item[0] for item in sorted_items]
        chart_data.add_series(mapping.get("series_name", "Data"), tuple([item[1] for item in sorted_items]))
        shape.chart.replace_data(chart_data)

    def _handle_aggregate_response_classes(self, slide_idx, shape, directive, response_data):
        extract_rules = directive.get("extract", {})
        dim_key = extract_rules.get("dimension", "responseClass")
        metric_count, metric_percent = extract_rules.get("metric_count", "edgeHitsSum"), extract_rules.get("metric_percent", "edgeHitsPercentage")
        
        class_data = {}
        for row in response_data.get("data", []):
            class_data[str(row.get(dim_key, "")).lower()] = {
                "count": float(row.get(metric_count, 0.0)),
                "percent": float(row.get(metric_percent, 0.0))
            }
            
        tag_values = {}
        for cls_key, tag_info in extract_rules.get("tags", {}).items():
            data = class_data.get(cls_key.lower(), {"count": 0.0, "percent": 0.0})
            tag_values[f"{{{{{tag_info.get('count')}}}}}"] = f"{int(data['count']):,}"
            tag_values[f"{{{{{tag_info.get('percent')}}}}}"] = f"{data['percent']:.2f} %"
            
        self._replace_tags_in_shape(shape, tag_values)

    def _extract_metric_value(self, data_rows, metric, operation, multiplier=1.0):
        values = [float(row.get(metric, 0)) for row in data_rows if metric in row]
        if not values: return 0.0
        if operation == "sum": return sum(values) * multiplier
        elif operation == "max": return max(values) * multiplier
        elif operation == "average": return (sum(values) / len(values)) * multiplier
        elif operation == "last": return values[-1] * multiplier
        return 0.0

    def _handle_replace_text(self, slide_idx, shape, directive, response_data):
        extract_rules = directive.get("extract", [])
        if not extract_rules: return
        data_rows, raw_values, tag_values = response_data.get("data", []), {}, {}
        
        for rule in extract_rules:
            op, target_tag, multiplier = rule.get("operation"), rule.get("target_tag"), rule.get("multiplier", 1.0)
            if op in ["sum_tags", "divide_tags"]: continue 
            
            if op == "max_sum_metrics":
                max_val = max([sum(float(row.get(m, 0)) for m in rule.get("metrics", []) if m in row) for row in data_rows] or [0.0])
                val = max_val * multiplier
            else:
                val = self._extract_metric_value(data_rows, rule.get("metric", ""), op, multiplier)
            raw_values[target_tag] = val
            tag_values[target_tag] = self._apply_unit_formatting(val, target_tag, rule.get("metric", ""))

        for rule in extract_rules:
            op, target_tag = rule.get("operation"), rule.get("target_tag")
            if op == "sum_tags":
                total_val = sum(raw_values.get(tag, 0.0) for tag in rule.get("source_tags", []))
                raw_values[target_tag] = total_val
                tag_values[target_tag] = self._apply_unit_formatting(total_val, target_tag)
            elif op == "divide_tags":
                num_input, den_input, multiplier = rule.get("numerator"), rule.get("denominator"), rule.get("multiplier", 1.0)
                num_val = sum(raw_values.get(t, 0.0) for t in num_input) if isinstance(num_input, list) else raw_values.get(num_input, 0.0)
                den_val = sum(raw_values.get(t, 0.0) for t in den_input) if isinstance(den_input, list) else raw_values.get(den_input, 0.0)
                result_val = (num_val / den_val * multiplier) if den_val != 0 else 0.0
                raw_values[target_tag] = result_val
                tag_values[target_tag] = self._apply_unit_formatting(result_val, target_tag)

        self._replace_tags_in_shape(shape, tag_values)
